"""Export ExpenseFlow monthly transactions to PDF and Excel."""
from collections import defaultdict
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle

from database import get_transactions


def monthly_data(year: int, month: int):
    """Return selected month's records and summary (independent of GUI filters)."""
    if not (1 <= month <= 12 and 2000 <= year <= 2100):
        raise ValueError('Choose a valid year and month.')
    prefix = f'{year:04d}-{month:02d}-'
    records = [r for r in get_transactions() if str(r[4]).startswith(prefix)]
    income = sum(float(r[3]) for r in records if r[1] == 'Income')
    expenses = sum(float(r[3]) for r in records if r[1] == 'Expense')
    return records, income, expenses, income - expenses


def export_pdf(destination, year, month):
    records, income, expenses, balance = monthly_data(year, month)
    destination = Path(destination)
    doc = SimpleDocTemplate(str(destination), pagesize=landscape(A4),
                            rightMargin=35, leftMargin=35, topMargin=32, bottomMargin=32)
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name='ReportTitle', parent=styles['Heading1'],
                              fontName='Helvetica-Bold', fontSize=20, textColor=colors.HexColor('#25205e'),
                              spaceAfter=8))
    styles.add(ParagraphStyle(name='ReportSmall', parent=styles['Normal'], fontSize=9, leading=12))
    story = [Paragraph('ExpenseFlow | Monthly Financial Report', styles['ReportTitle']),
             Paragraph(f'Report period: {year:04d}-{month:02d}  |  Generated: {date.today().isoformat()}', styles['ReportSmall']),
             Spacer(1, 16)]
    summary = [['TOTAL INCOME', 'TOTAL EXPENSES', 'NET BALANCE'],
               [f'Rs. {income:,.2f}', f'Rs. {expenses:,.2f}', f'Rs. {balance:,.2f}']]
    table = Table(summary, colWidths=[235, 235, 235])
    table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#202a44')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('BACKGROUND', (0, 1), (-1, 1), colors.HexColor('#edf1f9')),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 11),
        ('PADDING', (0, 0), (-1, -1), 12),
    ]))
    story.extend([table, Spacer(1, 22), Paragraph('Transactions', styles['Heading2']), Spacer(1, 8)])
    rows = [['ID', 'Type', 'Category', 'Amount (Rs.)', 'Date', 'Description']]
    for ident, kind, category, amount, when, description in records:
        # Paragraph handles long descriptions and escapes XML special characters.
        from xml.sax.saxutils import escape
        rows.append([str(ident), str(kind), str(category), f'{float(amount):,.2f}', str(when),
                     Paragraph(escape(str(description or '')), styles['ReportSmall'])])
    if not records:
        rows.append(['-', '-', '-', '-', '-', 'No transactions for this month'])
    history = Table(rows, colWidths=[42, 85, 105, 115, 90, 268], repeatRows=1, hAlign='LEFT')
    history.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#252f4a')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 9),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#f1f4fa')]),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 9),
        ('TOPPADDING', (0, 0), (-1, -1), 9),
        ('LINEBELOW', (0, 0), (-1, 0), 1, colors.HexColor('#6366f1')),
    ]))
    story.append(history)
    doc.build(story)
    return destination


def export_excel(destination, year, month):
    records, income, expenses, balance = monthly_data(year, month)
    destination = Path(destination)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Monthly Report'
    ws.merge_cells('A1:F1')
    ws['A1'] = 'ExpenseFlow | Monthly Financial Report'
    ws['A1'].font = Font(size=18, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill('solid', fgColor='202A44')
    ws['A1'].alignment = Alignment(vertical='center')
    ws.row_dimensions[1].height = 36
    ws['A2'] = 'Month'
    ws['B2'] = f'{year:04d}-{month:02d}'
    ws['A4'], ws['C4'], ws['E4'] = 'Total Income', 'Total Expenses', 'Net Balance'
    ws['A5'], ws['C5'], ws['E5'] = income, expenses, balance
    for cell in ('A4', 'C4', 'E4'):
        ws[cell].font = Font(bold=True, color='FFFFFF')
        ws[cell].fill = PatternFill('solid', fgColor='39466A')
    for cell in ('A5', 'C5', 'E5'):
        ws[cell].number_format = '"Rs. "#,##0.00;[Red]("Rs. "#,##0.00)'
        ws[cell].font = Font(bold=True, size=13)
    headings = ['ID', 'Type', 'Category', 'Amount (Rs.)', 'Date', 'Description']
    for col, label in enumerate(headings, 1):
        cell = ws.cell(8, col, label)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='252F4A')
    for row_index, (ident, kind, category, amount, when, description) in enumerate(records, 9):
        values = [ident, kind, category, float(amount), when, description or '']
        for col_index, value in enumerate(values, 1):
            cell = ws.cell(row_index, col_index, value)
            if row_index % 2 == 0:
                cell.fill = PatternFill('solid', fgColor='EEF2FA')
        ws.cell(row_index, 4).number_format = '#,##0.00'
    if not records:
        ws['A9'] = 'No transactions for this month'
    for col, width in {'A': 12, 'B': 17, 'C': 20, 'D': 21, 'E': 18, 'F': 44}.items():
        ws.column_dimensions[col].width = width
    ws.freeze_panes = 'A9'
    ws.auto_filter.ref = f'A8:F{max(9, 8 + len(records))}'
    wb.save(destination)
    return destination
